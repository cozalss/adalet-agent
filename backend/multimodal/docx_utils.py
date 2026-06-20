"""Word (.docx) metin çıkarımı — python-docx ile (pdf_utils'in Word karşılığı).

.docx zaten temiz bir metin katmanı taşır (XML) → OCR/OpenAI GEREKMEZ; doğrudan
paragraf + tablo metni çıkarılır (kusursuz Türkçe, bedava, hızlı). Yalnız modern
.docx; eski ikili .doc DESTEKLENMEZ (python-docx okuyamaz → çağıran 415/uyarı verir).
"""

from __future__ import annotations

import io
import os
import zipfile

# .docx = OOXML zip; içinde word/document.xml bulunur (düz zip/.xlsx'ten ayırt eder).
_ZIP_MAGIC = b"PK\x03\x04"
_DOCX_MARKER = "word/document.xml"
# Zip-bomb guard (CWE-409): 5 MB .docx açıldığında GB'lara şişebilir → worker OOM. Açılmış
# (uncompressed) toplam boyut bu tavanı aşarsa reddet. Meşru hukuk belgeleri << 200 MB.
_MAX_DOCX_UNCOMPRESSED = int(os.getenv("DOCX_MAX_UNCOMPRESSED_BYTES", str(200 * 1024 * 1024)))


def _guard_zip_bomb(data: bytes) -> None:
    """Açılmış toplam boyut tavanı aşılıyorsa ValueError (decompress ETMEDEN, header'dan)."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        total = sum(info.file_size for info in zf.infolist())
    if total > _MAX_DOCX_UNCOMPRESSED:
        raise ValueError(
            f".docx açılmış boyut sınırı aşıldı: {total} > {_MAX_DOCX_UNCOMPRESSED} bayt",
        )


def is_docx(data: bytes) -> bool:
    """Baytlar geçerli bir .docx mı? (zip magic + word/document.xml içeriği)."""
    if len(data) < 4 or data[:4] != _ZIP_MAGIC:
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            return _DOCX_MARKER in zf.namelist()
    except (zipfile.BadZipFile, OSError):
        return False


def extract_docx_text(data: bytes) -> str:
    """.docx'ten paragraf + tablo metnini çıkar (okuma sırasını yaklaşık korur).

    Bloklayıcı (zip/XML parse) → çağıran `asyncio.to_thread` + timeout ile sarmalı.
    """
    _guard_zip_bomb(data)  # OOM/zip-bomb koruması — docx.Document XML'i belleğe okur

    import docx  # python-docx (ana bağımlılık)

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in document.paragraphs if p.text and p.text.strip()]

    # Tablolar (hukuk belgelerinde sık: taraf/tutar/tarih çizelgeleri) — satır = "a | b | c".
    # dict.fromkeys: birleşik (merged) hücreler aynı metni tekrarlar → teksilleştir. c.text
    # bir property (her erişimde XML parse) → walrus ile bir kez oku.
    for table in document.tables:
        for row in table.rows:
            cells = list(dict.fromkeys(t for c in row.cells if (t := c.text.strip())))
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


__all__ = ("extract_docx_text", "is_docx")
