"""OCR — Word (.docx) çıkarımı + OpenAI vision routing + lokal fallback.

Gerçek ağ çağrısı YOK (ocr_image_openai / httpx.AsyncClient mock'lanır).
conftest autouse OPENAI_API_KEY'i temizler → default lokal (deterministik).
"""

from __future__ import annotations

import io

import pytest

import multimodal.ocr as ocr
import multimodal.ocr_cloud as ocr_cloud
from multimodal.docx_utils import _guard_zip_bomb, extract_docx_text, is_docx
from multimodal.ocr import extract_text
from multimodal.ocr_cloud import _detect_image_mime, _validated_base_url, cloud_ocr_enabled


def _make_png(size=(8, 8)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, "white").save(buf, format="PNG")
    return buf.getvalue()


def _make_docx(paragraphs, table_rows=None) -> bytes:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table_rows:
        t = d.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, cell in enumerate(row):
                t.rows[i].cells[j].text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
#  Word (.docx)
# --------------------------------------------------------------------------

def test_is_docx_true_for_real_docx():
    assert is_docx(_make_docx(["Merhaba"])) is True


def test_is_docx_false_for_pdf_and_random():
    assert is_docx(b"%PDF-1.7 ...") is False
    assert is_docx(b"not a zip at all") is False
    assert is_docx(b"PK\x03\x04randomzipnotdocx") is False  # zip ama docx değil


def test_extract_docx_text_paragraphs_and_table():
    data = _make_docx(
        ["Kira sözleşmesi", "Aylık kira 5000 TL"],
        table_rows=[["Taraf", "Ad"], ["Kiracı", "Ayşe Yılmaz"]],
    )
    text = extract_docx_text(data)
    assert "Kira sözleşmesi" in text
    assert "Aylık kira 5000 TL" in text
    assert "Kiracı | Ayşe Yılmaz" in text  # tablo satırı


@pytest.mark.asyncio
async def test_extract_text_routes_docx_to_docx_engine():
    data = _make_docx(["İşten haksız çıkarıldım, kıdem tazminatı isterim."])
    r = await extract_text(data)
    assert r.engine == "docx-text"
    assert "kıdem tazminatı" in r.text
    assert r.error is None


# --------------------------------------------------------------------------
#  OpenAI vision routing
# --------------------------------------------------------------------------

def test_use_cloud_ocr_routing(monkeypatch):
    monkeypatch.setattr(ocr, "OCR_PROVIDER", "local")
    monkeypatch.setenv("OPENAI_API_KEY", "sk-x")
    assert ocr._use_cloud_ocr() is False
    monkeypatch.setattr(ocr, "OCR_PROVIDER", "openai")
    assert ocr._use_cloud_ocr() is True
    monkeypatch.setattr(ocr, "OCR_PROVIDER", "auto")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    assert ocr._use_cloud_ocr() is False
    monkeypatch.setenv("OPENAI_API_KEY", "sk-x")
    assert ocr._use_cloud_ocr() is True


def test_cloud_enabled_reflects_key(monkeypatch):
    assert cloud_ocr_enabled() is False  # conftest temizledi
    monkeypatch.setenv("OPENAI_API_KEY", "sk-x")
    assert cloud_ocr_enabled() is True


@pytest.mark.asyncio
async def test_image_routes_to_cloud_vision(monkeypatch):
    monkeypatch.setattr(ocr, "OCR_PROVIDER", "openai")

    async def fake_vision(image_bytes, *, timeout_s, model=None):
        return "Madde 11 — ayıplı mal iadesi"

    monkeypatch.setattr(ocr, "ocr_image_openai", fake_vision)

    r = await extract_text(b"fake-jpeg-bytes")  # pdf/docx değil → _image_ocr → bulut
    assert r.engine == "openai-vision"
    assert "ayıplı mal" in r.text
    assert r.error is None


@pytest.mark.asyncio
async def test_cloud_failure_falls_back_to_local(monkeypatch):
    monkeypatch.setattr(ocr, "OCR_PROVIDER", "openai")
    monkeypatch.setattr(ocr, "OCR_CLOUD_FALLBACK", True)

    async def boom(image_bytes, *, timeout_s, model=None):
        raise RuntimeError("ağ hatası")

    monkeypatch.setattr(ocr, "ocr_image_openai", boom)

    async def fake_local(image_bytes, lang):
        return ocr.OCRResult(text="lokal rapidocr metni", engine="rapidocr",
                             lang=lang, n_chars=20, elapsed_ms=1)

    monkeypatch.setattr(ocr, "_rapidocr_ocr", fake_local)

    r = await extract_text(b"fake-jpeg-bytes")
    assert r.text == "lokal rapidocr metni"   # bulut patladı → lokal'e düştü
    assert r.engine == "rapidocr"


@pytest.mark.asyncio
async def test_cloud_failure_no_fallback_returns_error(monkeypatch):
    monkeypatch.setattr(ocr, "OCR_PROVIDER", "openai")
    monkeypatch.setattr(ocr, "OCR_CLOUD_FALLBACK", False)

    async def boom(image_bytes, *, timeout_s, model=None):
        raise RuntimeError("ağ hatası")

    monkeypatch.setattr(ocr, "ocr_image_openai", boom)

    r = await extract_text(b"fake-jpeg-bytes")
    assert r.text == ""
    assert r.error and "bulut" in r.error.lower()


# --------------------------------------------------------------------------
#  ocr_image_openai — istek kurulumu + MIME tespiti
# --------------------------------------------------------------------------

def test_detect_image_mime():
    assert _detect_image_mime(b"\xff\xd8\xff\xe0jpeg") == "image/jpeg"
    assert _detect_image_mime(b"\x89PNG\r\n\x1a\nrest") == "image/png"
    assert _detect_image_mime(b"RIFF????WEBPxx") == "image/webp"
    assert _detect_image_mime(b"unknownbytes") == "application/octet-stream"


@pytest.mark.asyncio
async def test_ocr_image_openai_builds_request(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-9")
    captured: dict = {}

    class FakeResp:
        def raise_for_status(self) -> None: pass
        def json(self) -> dict:
            return {"choices": [{"message": {"content": "  Madde 24  "}}]}

    class FakeClient:
        def __init__(self, *a, **k): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *a): return False
        async def post(self, url, headers=None, json=None):
            captured.update(url=url, headers=headers, json=json)
            return FakeResp()

    monkeypatch.setattr(ocr_cloud.httpx, "AsyncClient", FakeClient)

    text = await ocr_cloud.ocr_image_openai(_make_png(), timeout_s=30)

    assert text == "Madde 24"  # strip uygulandı
    assert captured["url"].endswith("/chat/completions")
    assert captured["headers"]["Authorization"] == "Bearer sk-test-9"
    content = captured["json"]["messages"][0]["content"]
    assert content[1]["image_url"]["url"].startswith("data:image/png;base64,")


@pytest.mark.asyncio
async def test_ocr_image_openai_no_key_raises():
    with pytest.raises(RuntimeError, match="OPENAI_API_KEY"):
        await ocr_cloud.ocr_image_openai(_make_png(), timeout_s=30)


# --------------------------------------------------------------------------
#  Güvenlik sertleştirme (ECC review 2026-06-20)
# --------------------------------------------------------------------------

def test_guard_zip_bomb_rejects_oversized(monkeypatch):
    """Açılmış boyutu sınırı aşan .docx reddedilmeli (zip-bomb / OOM koruması)."""
    monkeypatch.setattr("multimodal.docx_utils._MAX_DOCX_UNCOMPRESSED", 1000)
    big = _make_docx(["x" * 5000])  # açılmış document.xml > 1000 bayt
    with pytest.raises(ValueError, match="boyut sınırı"):
        _guard_zip_bomb(big)


def test_guard_zip_bomb_allows_normal():
    _guard_zip_bomb(_make_docx(["Kısa belge"]))  # raise etmemeli


def test_validated_base_url_rejects_unsafe(monkeypatch):
    monkeypatch.setenv("OPENAI_BASE_URL", "file:///etc/passwd")
    with pytest.raises(ValueError, match="güvensiz"):
        _validated_base_url()


def test_validated_base_url_allows_https(monkeypatch):
    monkeypatch.setenv("OPENAI_BASE_URL", "https://api.openai.com/v1")
    assert _validated_base_url() == "https://api.openai.com/v1"


@pytest.mark.asyncio
async def test_extract_text_rejects_docx_zip_bomb(monkeypatch):
    """extract_text zip-bomb .docx'i graceful hata ile döndürmeli (çökmeden)."""
    monkeypatch.setattr("multimodal.docx_utils._MAX_DOCX_UNCOMPRESSED", 1000)
    r = await extract_text(_make_docx(["x" * 5000]))
    assert r.engine == "docx-text"
    assert r.text == ""
    assert r.error  # "Word belgesi okunamadı" (ValueError yakalandı)
