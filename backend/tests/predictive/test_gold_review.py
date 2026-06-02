"""scripts/predictive/gold_review.py — CSV köprüsü saf çekirdek (export/ingest)."""

from __future__ import annotations

import csv

from scripts.predictive.gold_review import (
    export_csv,
    export_xlsx,
    ingest_csv,
    ingest_xlsx,
    normalize_label,
)

_CAND = [
    {"doc_id": "d1", "input_text": "Olay 1 tam metin", "hukum_excerpt": "h1",
     "qwen_label": "kazandi", "quote": "kabul", "needs_review": False},
    {"doc_id": "d2", "input_text": "Olay 2 tam metin", "hukum_excerpt": "h2",
     "qwen_label": "kaybetti", "quote": "ret", "needs_review": True},
]


def _fill_karar(csv_path, mapping: dict[str, str]) -> None:
    """CSV'nin 'karar' sütununu doc_id→değer ile doldur (hukukçuyu simüle et)."""
    rows = list(csv.DictReader(csv_path.read_text(encoding="utf-8-sig").splitlines()))
    for r in rows:
        r["karar"] = mapping.get(r["doc_id"], "")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def test_normalize_label_valid():
    assert normalize_label("kazandi") == "kazandi"
    assert normalize_label("  KAYBETTI ") == "kaybetti"
    assert normalize_label("Kısmi.") == "kismi"  # tr 'ı'→'i', nokta kırp


def test_normalize_label_invalid():
    assert normalize_label("") is None
    assert normalize_label("evet") is None


def test_export_writes_csv(tmp_path):
    csv_path = tmp_path / "review.csv"
    n = export_csv(_CAND, csv_path)
    assert n == 2 and csv_path.exists()
    header = csv_path.read_text(encoding="utf-8-sig").splitlines()[0]
    assert "karar" in header and "doc_id" in header


def test_export_then_ingest_roundtrip(tmp_path):
    csv_path = tmp_path / "review.csv"
    export_csv(_CAND, csv_path)
    _fill_karar(csv_path, {"d1": "kazandi", "d2": "kismi"})
    rows, skipped = ingest_csv(_CAND, csv_path)
    assert {r["label"] for r in rows} == {"kazandi", "kismi"}
    assert skipped == 0
    # full input_text join'lendi (CSV'deki kırpık değil)
    assert any(r["input_text"] == "Olay 1 tam metin" for r in rows)


def test_ingest_skips_empty_and_belirsiz(tmp_path):
    csv_path = tmp_path / "review.csv"
    export_csv(_CAND, csv_path)
    _fill_karar(csv_path, {"d1": "belirsiz"})  # d1 belirsiz (3-class dışı), d2 boş
    rows, skipped = ingest_csv(_CAND, csv_path)
    assert len(rows) == 0  # belirsiz eval'e girmez, boş atlanır
    assert skipped == 2


def test_xlsx_export_then_ingest_roundtrip(tmp_path):
    from openpyxl import load_workbook

    xlsx_path = tmp_path / "review.xlsx"
    assert export_xlsx(_CAND, xlsx_path) == 2
    assert xlsx_path.exists()
    # Hukukçuyu simüle et: KARAR (D sütunu) doldur — d1→kazandi, d2→kismi
    wb = load_workbook(xlsx_path)
    ws = wb.active
    ws["D2"], ws["D3"] = "kazandi", "kismi"
    wb.save(xlsx_path)

    rows, skipped = ingest_xlsx(_CAND, xlsx_path)
    assert {r["label"] for r in rows} == {"kazandi", "kismi"}
    assert skipped == 0
    assert any(r["input_text"] == "Olay 1 tam metin" for r in rows)
